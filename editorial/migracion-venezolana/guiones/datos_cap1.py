"""Reúne los datos del capítulo 1 en datos/capitulo-1.json.

Fuentes (carpeta insumos/, tal como llegaron en Capítulo 1_CODESS.zip):
  - Cuadros_Cap 1.xlsx ........ cuadros 1-9 y gráficas 7, 8 y 9
  - Cap1_Características.docx .. cuadros 4-8 (tablas de Word, 2022-2025)
  - Grafico2_*.html, Grafico3_*.html, Grafico4.html, Grafico5.html, mapa1.html, mapa2.html
  - Gráfica 1.png y Gráfica 6.png: solo existen como imagen; los valores se
    transcribieron de sus etiquetas (ver GRAFICA_1 y GRAFICA_6).

Uso: python guiones/datos_cap1.py
"""
import json
import pathlib
import re

import docx
import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
INS = RAIZ / "insumos"


def xlsx():
    return openpyxl.load_workbook(INS / "Cuadros_Cap 1.xlsx", data_only=True)


def num(v):
    """Números del Excel; algunos vienen como texto con separador '555, 484' o '32. 45'."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if re.fullmatch(r"\d{1,3}(,\d{3})+", s):
        return float(s.replace(",", ""))
    return float(s)


def filas(ws, desde, hasta, cols="BCDEFG"):
    out = []
    for r in range(desde, hasta + 1):
        etiqueta = ws[f"A{r}"].value
        if etiqueta is None and all(ws[f"{c}{r}"].value is None for c in cols):
            continue
        out.append({
            "etiqueta": re.sub(r"\s{2,}", " ", str(etiqueta or "")).strip(),
            "valores": [num(ws[f"{c}{r}"].value) for c in cols],
        })
    return out


def tablas_docx():
    d = docx.Document(INS / "Cap1_Características.docx")
    res = []
    for t in d.tables:
        filas_t = []
        for r in t.rows:
            celdas, previa = [], None
            for c in r.cells:
                if c._tc is previa:
                    continue
                previa = c._tc
                celdas.append(c.text.strip())
            filas_t.append(celdas)
        res.append(filas_t)
    return res


def piramide(nombre):
    txt = (INS / nombre).read_text(encoding="utf-8")
    bloque = re.search(r"const raw = \[(.*?)\];", txt, re.S).group(1)
    pares = re.findall(r"\[\s*(-?[\d.]+)\s*,\s*(-?[\d.]+)\s*\]", bloque)
    edades = ["0–4", "5–9", "10–14", "15–19", "20–24", "25–29", "30–34", "35–39", "40–44",
              "45–49", "50–54", "55–59", "60–64", "65–69", "70–74", "75–79", "80–84", "85+"]
    return [{"edad": e, "hombres": abs(float(h)), "mujeres": float(m)} for e, (h, m) in zip(edades, pares)]


# Gráfica 1 (solo PNG): millones de personas según fuente; None = sin dato
GRAFICA_1 = {
    "anios": list(range(2014, 2026)),
    "series": [
        {"clave": "mc", "nombre": "Registro administrativo (Migración Colombia)",
         "valores": [0.02, 0.03, 0.05, 0.4, 1.2, 1.8, 1.7, 1.8, 2.9, 2.9, 2.8, 2.8]},
        {"clave": "onu", "nombre": "ONU (International Migrant Stock)",
         "valores": [None, 0.03, None, None, None, None, 1.8, None, None, None, 2.9, None]},
        {"clave": "geih", "nombre": "GEIH",
         "valores": [None, None, None, None, None, None, 2.3, 2.8, 2.4, 2.3, 2.3, 2.1]},
        {"clave": "r4v", "nombre": "R4V",
         "valores": [None, None, None, None, 0.6, 1.2, 1.8, 1.7, 1.8, 2.5, 2.9, 2.8]},
    ],
}

# Gráfica 6 (solo PNG): % de la población venezolana ocupada, 2025.
# Hombres y mujeres, de las etiquetas del PNG. El PNG traía «Total (promedio)»
# = promedio simple de ambos sexos; aquí el total es el ponderado del cuadro 5.
GRAFICA_6 = [
    {"posicion": "Obrero o empleado de empresa particular", "hombres": 51.7, "mujeres": 49.9},
    {"posicion": "Trabajador por cuenta propia", "hombres": 43.4, "mujeres": 39.0},
    {"posicion": "Empleado doméstico", "hombres": 0.3, "mujeres": 7.3},
    {"posicion": "Trabajador familiar sin remuneración", "hombres": 0.6, "mujeres": 2.2},
    {"posicion": "Jornalero o peón", "hombres": 2.5, "mujeres": 0.2},
    {"posicion": "Patrón o empleador", "hombres": 1.5, "mujeres": 1.1},
    {"posicion": "Obrero o empleado del gobierno", "hombres": 0.0, "mujeres": 0.2},
    {"posicion": "Otro", "hombres": 0.0, "mujeres": 0.1},
]

# Nombres completos de las ramas (en el Excel vienen truncados a 21 caracteres)
RAMAS = {
    "Inmobiliaria": "Inmobiliaria",
    "Financieros y seguros": "Actividades financieras y de seguros",
    "Minería y canteras": "Minería y canteras",
    "Información y comunic": "Información y comunicaciones",
    "Electricidad, gas, ag": "Electricidad, gas y agua",
    "Administración públic": "Administración pública",
    "Trabajo doméstico en": "Trabajo doméstico en hogares",
    "Agropecuarío": "Agropecuario",
    "Profesionales y admin": "Profesionales y administrativas",
    "Transporte y almacena": "Transporte y almacenamiento",
    "Artísticas y otras ac": "Artísticas y otras actividades de servicios",
    "Construcción": "Construcción",
    "Industria manufacture": "Industria manufacturera",
    "Comercio y reparación": "Comercio y reparación",
    "Alojamiento y comida": "Alojamiento y servicios de comida",
}


def main():
    wb = xlsx()
    tdocx = tablas_docx()
    datos = {}

    # --- Mapa 1 (mapa1.html) ---
    m1 = (INS / "mapa1.html").read_text(encoding="utf-8")
    datos["mapa_1"] = [
        {"pais": n, "valor": int(v)}
        for n, v in re.findall(r'\{\s*name:\s*"([^"]+)",\s*value:\s*(\d+)\s*\}', m1)
    ]

    # --- Mapa 2 (mapa2.html) ---
    m2 = (INS / "mapa2.html").read_text(encoding="utf-8")
    def dic(nombre):
        b = re.search(rf"const {nombre} = \{{(.*?)\}};", m2, re.S).group(1)
        return {k: float(v) for k, v in re.findall(r'"([^"]+)":\s*([\d.]+)', b)}
    datos["mapa_2"] = {"total": dic("total"), "sin_documentos": dic("irregular")}

    # --- Diagrama 1 (Diagrama 1.jpg) ---
    datos["diagrama_1"] = [
        {"periodo": "1965-1999", "etapa": "Periodo previo a la crisis", "valor": 0.65},
        {"periodo": "2000-2012", "etapa": "Inicia emigración de manera selectiva", "valor": 2.33},
        {"periodo": "2013-2015", "etapa": "Crisis aguda", "valor": 5.37},
        {"periodo": "2016-2022", "etapa": "Migración masiva", "valor": 76.76},
        {"periodo": "2023-2025", "etapa": "Periodo poscrisis", "valor": 14.83},
    ]

    datos["grafica_1"] = GRAFICA_1

    # --- Gráficas 2 y 3: pirámides (Grafico3_2022.html es la población colombiana 2025) ---
    datos["grafica_2"] = {"2022": piramide("Grafico2_2022.html"), "2025": piramide("Grafico2_2025.html")}
    datos["grafica_3"] = {"otros": piramide("Grafico3_2025.html"), "colombiana": piramide("Grafico3_2022.html")}

    # --- Gráfica 4 (Grafico4.html): M = hombres, F = mujeres ---
    g4 = (INS / "Grafico4.html").read_text(encoding="utf-8")
    datos["grafica_4"] = [
        {"nivel": n, "total": float(t), "hombres": float(h), "mujeres": float(m)}
        for n, t, h, m in re.findall(
            r'level:\s*"([^"]+)",\s*total:\s*([\d.]+),\s*M:\s*([\d.]+),\s*F:\s*([\d.]+)', g4)
    ]

    # --- Gráfica 5 (Grafico5.html) ---
    g5 = (INS / "Grafico5.html").read_text(encoding="utf-8")
    datos["grafica_5"] = [
        {"id": i, "nombre": n,
         "mujeres": [float(x) for x in mu.split(",")],
         "hombres": [float(x) for x in ho.split(",")]}
        for i, n, mu, ho in re.findall(
            r'id:\s*"(\w+)".*?nombre:\s*"([^"]+)".*?mujer:\s*\[([^\]]+)\],\s*hombre:\s*\[([^\]]+)\]', g5)
    ]

    # --- Gráfica 6 ---
    datos["grafica_6"] = GRAFICA_6

    # --- Gráficas 7 y 8 (hoja «Gráfica 7»: B = hombres, C = mujeres, D = total) ---
    ws = wb["Gráfica 7"]
    ramas = []
    for r in range(3, 18):
        corto = ws[f"A{r}"].value
        ramas.append({"rama": RAMAS[corto], "rama_excel": corto, "hombres": num(ws[f"B{r}"].value),
                      "mujeres": num(ws[f"C{r}"].value), "total": num(ws[f"D{r}"].value)})
    datos["grafica_7_8"] = sorted(ramas, key=lambda x: -x["total"])

    # --- Gráfica 9 (hoja «Gráfica 9», pesos -> miles de pesos) ---
    ws = wb["Gráfica 9"]
    g9 = {}
    for grupo, r0 in (("Venezolana", 16), ("Colombiana", 20), ("Otra población migrante", 24)):
        g9[grupo] = {
            cond: {"hombres": num(ws[f"B{r}"].value) / 1000, "mujeres": num(ws[f"C{r}"].value) / 1000}
            for cond, r in (("Formal", r0), ("Informal", r0 + 1), ("Total población ocupada", r0 + 2))
        }
    datos["grafica_9"] = g9

    # --- Cuadros 1-3 (Excel) ---
    datos["cuadro_1"] = filas(wb["Cuadro 1"], 3, 20)
    datos["cuadro_2"] = filas(wb["Cuadro 2"], 3, 19)
    datos["cuadro_3"] = filas(wb["Cuadro 3"], 3, 28)

    # --- Cuadros 4-8: tablas del DOCX (el cuadro 4 del Excel solo tiene 2025) ---
    def tabla(i):
        return [[c for c in fila] for fila in tdocx[i]]
    datos["cuadro_4"] = tabla(2)
    datos["cuadro_5"] = tabla(4)
    datos["cuadro_6"] = tabla(5)
    datos["cuadro_7"] = tabla(6)
    datos["cuadro_8"] = tabla(8)

    # --- Cuadro 9 (Excel, proporciones -> %) ---
    ws = wb["Cuadro 9"]
    datos["cuadro_9"] = [
        {"grupo": g, "total": num(ws[f"B{r}"].value) * 100,
         "informal": num(ws[f"C{r}"].value) * 100, "formal": num(ws[f"D{r}"].value) * 100}
        for g, r in (("Población venezolana", 3), ("Otros migrantes", 5), ("Población colombiana", 7))
    ]

    destino = RAIZ / "datos" / "capitulo-1.json"
    destino.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    print("->", destino.relative_to(RAIZ), ", ".join(datos))


if __name__ == "__main__":
    main()
